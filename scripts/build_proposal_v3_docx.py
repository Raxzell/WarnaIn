"""
WarnaIn v3 - Proposal Document Generator
Generates:
  - keperluan submisi/Proposal_WarnaIn_v3.docx
  - WarnaIn3/reports/Proposal_WarnaIn_v3.docx
Aligned with 3 Deep Learning Paradigms (CNN vs GAN vs Transformer) and LO3/LO4 objectives.
"""

import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2.5)
    s.left_margin = Cm(3)
    s.right_margin = Cm(2.5)

st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.line_spacing = 1.5

for name, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12)):
    h = doc.styles[name]
    h.font.name = "Times New Roman"
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    h.font.size = Pt(size)
    h.font.bold = True
    h.font.color.rgb = RGBColor(0, 0, 0)


def para(text, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=None):
    p = doc.add_paragraph()
    p.alignment = align
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    return p


def rich(label, text, style=None):
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(label).bold = True
    p.add_run(text)
    return p


def bullets(items):
    for it in items:
        if isinstance(it, tuple):
            rich(it[0], it[1], "List Bullet")
        else:
            p = doc.add_paragraph(it, style="List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def numbered(items):
    for it in items:
        if isinstance(it, tuple):
            rich(it[0], it[1], "List Number")
        else:
            p = doc.add_paragraph(it, style="List Number")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(10.5)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade(c, "D9D9D9")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            cells[i].paragraphs[0].add_run(v).font.size = Pt(10.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.line_spacing = 1.15
    doc.add_paragraph()
    return t


# ---------------- Cover ----------------
for _ in range(3):
    doc.add_paragraph()
para("PROPOSAL PROYEK SISTEM CERDAS & COMPUTER VISION", True, align=WD_ALIGN_PARAGRAPH.CENTER, size=15)
para("WarnaIn v3: Komparasi Kinerja dan Analisis Perilaku Tiga Paradigma Deep Learning "
     "(CNN vs. GAN vs. Vision Transformer) pada Restorasi Citra Monokrom "
     "Berbasis Dekomposisi Ruang Warna CIE L*a*b*",
     True, align=WD_ALIGN_PARAGRAPH.CENTER, size=13)
for _ in range(5):
    doc.add_paragraph()
para("Disusun oleh:", align=WD_ALIGN_PARAGRAPH.CENTER)
para("Ihsan Panji Rahmawan\nJames Michael Lionel\nMuhammad Rasyadu", True, align=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(2):
    doc.add_paragraph()
para("Program Studi Computer Science\nFakultas Ilmu Komputer\n2026", align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

# ---------------- A. Latar Belakang ----------------
doc.add_heading("A. Latar Belakang", 1)
numbered([
    ("Tuntutan Capaian Pembelajaran Akademik (LO3 & LO4). ",
     "Sesuai arahan dewan penguji akademik, pengembangan sistem restorasi citra tidak hanya dituntut untuk "
     "membangun artefak aplikasi yang berfungsi, melainkan wajib membuktikan pencapaian Learning Outcome 3 "
     "(LO3: Analisis Perilaku Model) dan Learning Outcome 4 (LO4: Evaluasi Limitasi Pendekatan Deep Learning). "
     "Oleh karena itu, proyek 'WarnaIn v3' direfaktor secara komprehensif untuk mengomparasikan secara langsung "
     "tiga keluarga model yang merepresentasikan tiga paradigma pembelajaran deep learning yang secara fundamental "
     "berbeda: (1) Regresi Piksel Langsung (CNN ResNet-18 U-Net), (2) Pembelajaran Adversarial (Conditional GAN Pix2Pix), "
     "dan (3) Mekanisme Self-Attention Global (Vision Transformer Swin-UNet)."),
    ("Tantangan Ambiguitas Multimodal pada Ruang Warna RGB. ",
     "Restorasi pewarnaan citra grayscale merupakan masalah ill-posed dengan sifat ambiguitas multimodal tinggi: "
     "sebuah objek abu-abu (misalnya pakaian, kendaraan, atau fasad bangunan) memiliki banyak warna kromatis yang "
     "sama-sama valid secara visual. Pada ruang warna RGB konvensional, intensitas cahaya dan rona warna terikat erat "
     "pada ketiga kanal. Apabila jaringan dipaksa memetakan citra grayscale langsung ke RGB, optimasi numerik "
     "cenderung mengalami gejala regression-to-the-mean yang menghasilkan citra kusam/desaturasi kecokelatan serta "
     "mengaburkan detail tepi (blurring)."),
    ("Keunggulan Dekomposisi Ruang Warna CIE L*a*b*. ",
     "WarnaIn memformulasikan pewarnaan sebagai pemetaan ortogonal pada ruang warna CIE L*a*b*. Kanal luminansi L* "
     "(rentang [-1, 1]) memuat 100% informasi iluminasi, gradien tekstur, dan kontur geometris citra asli. Dengan "
     "mempertahankan kanal L* sumber tanpa modifikasi (lossless) dan hanya menugaskan model neural untuk memprediksi "
     "dua matriks krominansi a*b* (rentang [-1, 1]), ketajaman struktural citra dijamin tetap utuh tanpa degradasi spasial."),
    ("Relevansi Komparasi Tiga Paradigma Arsitektur. ",
     "Ketiga arsitektur yang diangkat mencerminkan filosofi komputasi yang kontras: "
     "(a) CNN (ResNet-18 U-Net) mengekstraksi fitur lokal bertingkat melalui operasi konvolusi dan skip connections; "
     "(b) GAN (Pix2Pix cGAN) memanfaatkan interaksi minimax antara U-Net Generator dan PatchGAN 70x70 Discriminator "
     "untuk memaksa keluaran mendekati manifold visual alami tanpa regresi rata-rata; dan "
     "(c) Vision Transformer (Swin-UNet) mengadopsi Windowed Multi-Head Self-Attention (W-MSA & SW-MSA) untuk memodelkan "
     "korelasi semantik jarak jauh di seluruh penjuru citra."),
    ("Inovasi Studio Web Editorial Broadsheet WarnaIn v3. ",
     "Seluruh arsitektur diintegrasikan ke dalam Web Studio interaktif berbasis Streamlit dengan antarmuka surat kabar "
     "editorial klasik. WarnaIn v3 menghadirkan fitur baru 'Head-to-Head 3-Model Comparison' (4 kolom berdampingan "
     "menampilkan input L*, output CNN, GAN, dan Transformer beserta metrik instan), CIE Lab Space Inspector, dan "
     "pengujian degradasi saturasi (0–100%) untuk membuktikan ketahanan model."),
])

# ---------------- B. Rumusan Masalah & Tujuan ----------------
doc.add_heading("B. Rumusan Masalah dan Tujuan", 1)
doc.add_heading("1. Rumusan Masalah", 2)
numbered([
    "Bagaimana merancang struktur model modular (Multi-Architecture Factory) yang membungkus CNN (ResNet-18 U-Net), "
    "Adversarial cGAN (Pix2Pix), dan Vision Transformer (Swin-UNet) dengan antarmuka I/O standar berdimensi (B, 1, 224, 224) "
    "menjadi (B, 2, 224, 224) dalam rentang [-1, 1]?",
    "Bagaimana karakteristik perilaku visual (LO3) masing-masing paradigma dalam mensintesis krominansi a*b* "
    "terhadap batas kontur, saturasi warna, dan konsistensi semantik?",
    "Apa saja limitasi matematis dan komputasional (LO4) dari pendekatan regresi piksel (Smooth L1/MSE), adversarial loss (PatchGAN), "
    "dan self-attention (W-MSA) terkait fenomena regression-to-the-mean, artefak halusinasi, dan latensi inferensi?",
    "Bagaimana membangun antarmuka web interaktif yang memungkinkan pengujian komparatif head-to-head secara simultan "
    "dan reproduktif pada dataset benchmark COCO 2017?",
])

doc.add_heading("2. Tujuan Proyek", 2)
numbered([
    ("Tujuan Rekayasa Arsitektur: ", "Mengimplementasikan Multi-Architecture Factory di src/model.py yang mencakup "
     "ResNetUNetColorizer, Pix2PixColorizer (Generator U-Net + Discriminator PatchGAN 70x70), dan TransformerColorizer "
     "(Swin-UNet / Restormer) berbasis PyTorch native."),
    ("Tujuan Evaluasi Empiris (LO3 & LO4): ", "Melakukan benchmark kuantitatif otomatis pada 500 citra uji COCO 2017 "
     "menggunakan metrik PSNR, SSIM, MAE kanal a*b*, dan Latensi per frame (ms), serta mengekspornya ke format CSV dan Markdown."),
    ("Tujuan Pengembangan Platform: ", "Memperbarui antarmuka web editorial Streamlit (app/app.py) dengan fitur Head-to-Head "
     "4 Kolom, kartu analisis karakteristik paradigma, slider Before-After interaktif, dan inspektor CIE L*a*b*."),
    ("Tujuan Reproduksibilitas: ", "Menyediakan notebook pelatihan terintegrasi (WarnaIn_v3_Training_Benchmark.ipynb) "
     "yang dapat dieksekusi secara modular baik pada lingkungan komputasi lokal maupun cloud GPU (Google Colab/Kaggle)."),
])

# ---------------- C. Metodologi Penelitian & Pengembangan ----------------
doc.add_heading("C. Metodologi Penelitian dan Pengembangan", 1)
para("Pengembangan WarnaIn v3 dilaksanakan mengikuti metodologi rekayasa perangkat lunak berbasis eksperimen terstruktur:")

doc.add_heading("1. Standarisasi Pipeline Data dan Ruang Warna", 2)
bullets([
    ("Dataset Acuan: ", "Dataset Microsoft COCO 2017 (5.000 citra validasi) yang telah dibagi secara seimbang "
     "menjadi train (4.000 citra), val (500 citra), dan test (500 citra) dengan seed 42, terorganisir ke dalam 5 kategori "
     "semantik (animals, architecture, landscapes, portraits, objects) serta 1 mode general."),
    ("Transformasi CIE L*a*b* [-1, 1]: ", "Citra RGB dikonversi ke CIE L*a*b* presisi float32. Kanal L* [0, 100] dinormalisasi "
     "ke [-1, 1] melalui formulasi L_norm = (L / 50.0) - 1.0. Kanal a*b* [-128, 127] dinormalisasi ke [-1, 1] melalui ab_norm = ab / 128.0."),
    ("Tensor Input-Output: ", "Input generator berdimensi (B, 1, 224, 224) bertipe float32 dalam rentang [-1, 1], dan output "
     "berdimensi (B, 2, 224, 224) dalam rentang [-1, 1] dengan fungsi aktivasi akhir Tanh."),
])

doc.add_heading("2. Desain Tiga Paradigma Arsitektur Deep Learning", 2)
bullets([
    ("Paradigma 1 — CNN: ResNet-18 U-Net (Pixel Regression): ",
     "Encoder mengadaptasi backbone ResNet-18 pra-latih ImageNet dengan memodifikasi conv1 menjadi 1-channel (bobot dirata-ratakan). "
     "Decoder terdiri dari 5 tahap progressive transposed convolution (ConvTranspose2d 4x4, stride 2) yang dihubungkan dengan "
     "skip connections dari stem, layer1, layer2, layer3, dan layer4."),
    ("Paradigma 2 — Adversarial: Pix2Pix cGAN (Conditional GAN): ",
     "Generator U-Net 5-stage downsampling (LeakyReLU 0.2) dan upsampling dengan Dropout 0.5 pada lapisan dalam untuk mencegah kolaps mode. "
     "Discriminator mengimplementasikan PatchGAN 70x70 yang menerima 3-channel masukan [L*, a*b*] (konkat L asli dengan a*b* nyata/palsu) "
     "untuk membedakan realisme patch beresolusi 70x70."),
    ("Paradigma 3 — Transformer: Swin-UNet (Windowed Multi-Head Self-Attention): ",
     "Arsitektur hierarkis Vision Transformer murni PyTorch berbasis blok Swin Transformer. Menggunakan Window Attention 2D "
     "(window size 7x7) dengan kombinasi regular windowing (W-MSA) dan shifted windowing (SW-MSA) berselang-seling, "
     "dilengkapi relative position bias table dan Gated Depthwise Feed-Forward Network (GDFN)."),
])

doc.add_heading("3. Fungsi Objektif dan Strategi Pelatihan", 2)
bullets([
    ("Fungsi Loss CNN (Model 1): ", "Kombinasi regresi L_total = 0.7 * SmoothL1(ab_pred, ab_gt) + 0.3 * MSE(ab_pred, ab_gt)."),
    ("Fungsi Loss GAN (Model 2): ", "Optimasi minimax bolak-balik: L_D = 0.5 * [BCE(D(L, ab_real), 1) + BCE(D(L, ab_fake), 0)] "
     "dan L_G = BCE(D(L, ab_fake), 1) + 100 * L1(ab_fake, ab_gt)."),
    ("Fungsi Loss Transformer (Model 3): ", "Smooth L1 Loss dioptimasi menggunakan optimizer AdamW dengan Cosine Annealing Learning Rate."),
])

doc.add_heading("4. Lingkungan dan Protokol Pengujian Komparatif", 2)
para("Pengujian dilakukan secara kuantitatif pada seluruh 500 citra uji COCO 2017 test set menggunakan skrip src/evaluate.py. "
     "Metrik yang diukur meliputi:")
bullets([
    ("PSNR (Peak Signal-to-Noise Ratio): ", "Mengukur fidelitas rekonstruksi sinyal dalam satuan desibel (dB) terhadap ground-truth RGB."),
    ("SSIM (Structural Similarity Index): ", "Mengukur kemiripan perseptual struktur, kontras, dan luminansi (rentang 0–1)."),
    ("MAE (Mean Absolute Error pada kanal a*b*): ", "Mengukur rata-rata penyimpangan absolut pada domain krominansi CIE a*b* [0, 128]."),
    ("Latency (Waktu Inferensi per frame): ", "Mengukur waktu komputasi forward-pass murni dalam milidetik (ms) pada resolusi 224x224."),
])

# ---------------- D. Rencana Implementasi & Jadwal Kerja ----------------
doc.add_heading("D. Rencana Implementasi dan Jadwal Kerja", 1)
table(
    ["Tahap", "Kegiatan Utama", "Output / Target", "Alokasi"],
    [
        ["Tahap I", "Refactoring model.py (Multi-Architecture Factory 3 Paradigma)", "Modul model.py terverifikasi dengan input-output [-1, 1]", "Minggu 1"],
        ["Tahap II", "Pembuatan skrip benchmark evaluate.py & standarisasi dataset", "Skrip otomatisasi metrik PSNR, SSIM, MAE, dan Latency", "Minggu 2"],
        ["Tahap III", "Penyusunan notebook training WarnaIn_v3_Training_Benchmark.ipynb", "Pipeline pelatihan 3 model siap jalan di lokal/Colab", "Minggu 2-3"],
        ["Tahap IV", "Pembaruan Streamlit UI (app.py) dengan Head-to-Head 4 Kolom", "Antarmuka web interaktif v3 editorial broadsheet", "Minggu 3"],
        ["Tahap V", "Eksekusi benchmark 500 citra COCO & penyusunan laporan Bab 4", "benchmark_summary.csv & laporan komparasi LO3/LO4", "Minggu 4"],
    ],
    widths=[2.5, 6.0, 5.0, 2.5]
)

# ---------------- E. Hasil yang Diharapkan ----------------
doc.add_heading("E. Hasil yang Diharapkan dan Kriteria Keberhasilan", 1)
numbered([
    ("Tersedianya Artefak Perangkat Lunak Teruji: ", "Aplikasi web WarnaIn v3 yang stabil, modular, dan responsif, "
     "mampu menjalankan inferensi pada ketiga paradigma secara bergantian maupun bersamaan dalam mode komparasi 4 kolom."),
    ("Pencapaian Benchmark Kuantitatif: ", "Rata-rata SSIM > 0.90 pada test set 500 citra COCO untuk membuktikan bahwa preservasi "
     "kanal L* berhasil mempertahankan kontur citra, serta PSNR > 22 dB untuk seluruh arsitektur."),
    ("Pembuktian Pembelajaran Akademik (LO3 & LO4): ", "Dokumentasi komparatif yang menyajikan bukti empiris tentang fenomena "
     "regression-to-the-mean pada CNN, trade-off antara vibransi warna vs artefak pada GAN, serta keunggulan koherensi global "
     "berbiaya latensi lebih tinggi pada Vision Transformer."),
])

# ---------------- F. Referensi ----------------
doc.add_heading("F. Referensi", 1)
refs = [
    "[1] R. Zhang, P. Isola, and A. A. Efros, \"Colorful Image Colorization,\" in European Conference on Computer Vision (ECCV), 2016, pp. 649–666.",
    "[2] P. Isola, J.-Y. Zhu, T. Zhou, and A. A. Efros, \"Image-to-Image Translation with Conditional Adversarial Networks,\" in IEEE Conference on Computer Vision and Pattern Recognition (CVPR), 2017, pp. 1125–1134.",
    "[3] Z. Liu et al., \"Swin Transformer: Hierarchical Vision Transformer using Shifted Windows,\" in IEEE/CVF International Conference on Computer Vision (ICCV), 2021, pp. 10012–10022.",
    "[4] H. Cao et al., \"Swin-Unet: Unet-like Pure Transformer for Medical Image Segmentation,\" in European Conference on Computer Vision (ECCV) Workshops, 2022, pp. 205–218.",
    "[5] O. Ronneberger, P. Fischer, and T. Brox, \"U-Net: Convolutional Networks for Biomedical Image Segmentation,\" in MICCAI, 2015, pp. 234–241.",
    "[6] K. He, X. Zhang, S. Ren, and J. Sun, \"Deep Residual Learning for Image Recognition,\" in IEEE CVPR, 2016, pp. 770–778.",
    "[7] T.-Y. Lin et al., \"Microsoft COCO: Common Objects in Context,\" in European Conference on Computer Vision (ECCV), 2014, pp. 740–755.",
    "[8] Z. Wang, A. C. Bovik, H. R. Sheikh, and E. P. Simoncelli, \"Image Quality Assessment: From Error Visibility to Structural Similarity,\" IEEE Trans. Image Process., vol. 13, no. 4, pp. 600–612, 2004.",
    "[9] Y. Blau and T. Michaeli, \"The Perception-Distortion Tradeoff,\" in IEEE CVPR, 2018, pp. 6228–6237.",
]
for r in refs:
    p = doc.add_paragraph(r)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15

# Save to destination paths
out_submission = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "keperluan submisi", "Proposal_WarnaIn_v3.docx"))
out_reports = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "WarnaIn3", "reports", "Proposal_WarnaIn_v3.docx"))
os.makedirs(os.path.dirname(out_submission), exist_ok=True)
os.makedirs(os.path.dirname(out_reports), exist_ok=True)

doc.save(out_submission)
doc.save(out_reports)
print(f"[SUCCESS] Proposal v3 saved to:\n  - {out_submission}\n  - {out_reports}")
